"""The management view: every item, by month, company, status and approver, with Excel export.

JSAP's equivalent was a month-wise count per user, a DocEntry / vendor search
and a three-sheet export (`budget.js`). Here it is one filter set over
`BudgetItem`, read by both the page and the export, so the two never disagree.

The month is the DOCUMENT's month (the draft's date in SAP) — the month the
spend belongs to — not when OMS happened to take it in.

Read-only. Nobody decides anything from here.
"""
import io
from collections import defaultdict
from datetime import date

from django.contrib.auth import get_user_model
from django.db.models import Count, Q, Sum

from budget.models import BudgetActionLog, BudgetItem, ItemStatus, LogAction

#: Rows a page or an export carries at most.
LIMIT = 5000


class ReportInvalid(ValueError):
    pass


def _month(value):
    """`YYYY-MM` -> (first day, first day of the next month), or None."""
    value = (value or '').strip()
    if not value:
        return None
    try:
        year, month = (int(p) for p in value.split('-'))
        start = date(year, month, 1)
    except (TypeError, ValueError):
        raise ReportInvalid('`month` must look like 2026-10.') from None
    return start, date(year + (month == 12), month % 12 + 1, 1)


def items(params):
    """The filtered items. `params`: month, company, status, q (DocEntry / DocNum, vendor, budget head)."""
    qs = (BudgetItem.objects.select_related('draft', 'workflow', 'current_stage', 'current_user')
          .prefetch_related('lines'))
    span = _month(params.get('month'))
    if span:
        qs = qs.filter(draft__doc_date__gte=span[0], draft__doc_date__lt=span[1])
    company = (params.get('company') or '').strip().upper()
    if company:
        qs = qs.filter(company=company)
    status = (params.get('status') or '').strip().upper()
    if status:
        if status not in ItemStatus.values:
            raise ReportInvalid('`status` must be one of ' + ', '.join(ItemStatus.values) + '.')
        qs = qs.filter(status=status)
    q = (params.get('q') or '').strip()
    if q:
        match = Q(draft__card_name__icontains=q) | Q(budget_code__icontains=q) | Q(draft__card_code__iexact=q)
        if q.isdigit():
            match |= Q(draft__draft_entry=int(q)) | Q(draft__doc_num=int(q))
        qs = qs.filter(match)
    return qs.order_by('-draft__doc_date', '-id')


def summary(qs):
    """Counts and amounts per status, over the filtered items."""
    rows = qs.order_by().values('status').annotate(n=Count('id'), amount=Sum('amount'))
    by = {r['status']: r for r in rows}
    out = {s.lower(): {'count': by.get(s, {}).get('n', 0), 'amount': str(by.get(s, {}).get('amount') or 0)}
           for s in ItemStatus.values}
    out['total'] = {'count': sum(r['n'] for r in rows), 'amount': str(sum((r['amount'] or 0) for r in rows))}
    return out


def approvers(qs):
    """Per person: waiting at their stage now, and what they approved / rejected among these items.

    Auto-approvals are counted against the stage's configured user they stood
    in for, as JSAP's report did — the person who did not get to it.
    """
    from workflow.models import WorkflowStage

    people = defaultdict(lambda: {'pending': 0, 'approved': 0, 'auto_approved': 0, 'rejected': 0})
    names = {}
    ids = list(qs.values_list('id', flat=True)[:LIMIT])
    for row in (qs.filter(status=ItemStatus.PENDING).exclude(current_user=None).order_by()
                .values('current_user', 'current_user__name', 'current_user__username')
                .annotate(n=Count('id'))):
        people[row['current_user']]['pending'] = row['n']
        names[row['current_user']] = row['current_user__name'] or row['current_user__username']
    logs = (BudgetActionLog.objects.filter(item_id__in=ids, action__in=[LogAction.APPROVE, LogAction.REJECT,
                                                                        LogAction.AUTO_APPROVE])
            .values('action', 'acted_by', 'acted_by__name', 'acted_by__username', 'stage_id'))
    stage_owner = dict(WorkflowStage.objects.filter(id__in={r['stage_id'] for r in logs if r['stage_id']})
                       .values_list('id', 'user_id'))
    for u in get_user_model().objects.filter(pk__in=set(stage_owner.values())):
        names.setdefault(u.pk, u.name or u.username)
    for r in logs:
        if r['action'] == LogAction.AUTO_APPROVE:
            who = stage_owner.get(r['stage_id'])
            key = 'auto_approved'
        else:
            who = r['acted_by']
            key = 'approved' if r['action'] == LogAction.APPROVE else 'rejected'
            if who:
                names.setdefault(who, r['acted_by__name'] or r['acted_by__username'])
        if who:
            people[who][key] += 1
    out = [{'user_id': uid, 'name': names.get(uid, f'User {uid}'), **counts,
            'total': sum(counts.values())} for uid, counts in people.items()]
    return sorted(out, key=lambda r: (-r['pending'], r['name'].lower()))


def _decided(item):
    """(who, when, remarks) of the decision that closed it, else blanks."""
    log = None
    for row in item.action_logs.all():
        if row.action in (LogAction.APPROVE, LogAction.REJECT, LogAction.AUTO_APPROVE):
            log = row
    if log is None or item.status == ItemStatus.PENDING:
        return '', None, ''
    who = 'Auto-approval' if log.action == LogAction.AUTO_APPROVE else (
        (log.acted_by.name or log.acted_by.username) if log.acted_by else '')
    return who, log.acted_at, log.remarks


def export_xlsx(qs, params):
    """Three sheets — Approvers, Items, Lines — as `.xlsx` bytes."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    rows = list(qs.prefetch_related('action_logs__acted_by')[:LIMIT])
    book = Workbook()
    bold = Font(bold=True)

    def sheet(ws, header, data):
        ws.append(header)
        for cell in ws[1]:
            cell.font = bold
        for line in data:
            ws.append(line)
        for col in ws.columns:
            width = max(len(str(c.value or '')) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 60)
        ws.freeze_panes = 'A2'

    ws = book.active
    ws.title = 'Approvers'
    sheet(ws, ['Approver', 'Pending', 'Approved', 'Auto-approved', 'Rejected', 'Total'],
          [[a['name'], a['pending'], a['approved'], a['auto_approved'], a['rejected'], a['total']]
           for a in approvers(qs)])

    def decided_cols(item):
        who, when, remarks = _decided(item)
        return [who, when.replace(tzinfo=None) if when else None, remarks]

    sheet(book.create_sheet('Items'),
          ['Company', 'Document', 'Draft', 'Doc date', 'Party', 'Raised in SAP by', 'Budget head', 'Amount',
           'Status', 'Waiting at', 'Waiting on', 'Waiting since', 'Decided by', 'Decided at', 'Remarks',
           'SAP gate'],
          [[i.company, f'{i.draft.get_obj_type_display()} {i.draft.doc_num or i.draft.draft_entry}',
            i.draft.draft_entry, i.draft.doc_date, i.draft.card_name, i.draft.sap_created_by, i.budget_code,
            float(i.amount), i.get_status_display(),
            i.current_stage.name if i.current_stage_id and i.current_stage else '',
            (i.current_user.name or i.current_user.username) if i.current_user else '',
            i.waiting_since.replace(tzinfo=None) if i.waiting_since and i.status == ItemStatus.PENDING else None,
            *decided_cols(i), i.sap_status or ''] for i in rows])
    sheet(book.create_sheet('Lines'),
          ['Company', 'Draft', 'Line', 'Account', 'Account name', 'Budget head', 'Sub-budget', 'Effect month',
           'Amount', 'Line remarks', 'Item status'],
          [[i.company, i.draft.draft_entry, ln.line_num, ln.acct_code, ln.acct_name, ln.budget_code,
            ln.sub_budget_code, ln.effect_month, float(ln.amount), ln.remarks, i.get_status_display()]
           for i in rows for ln in i.lines.all()])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
